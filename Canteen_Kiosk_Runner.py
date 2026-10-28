# -*- coding: utf-8 -*-
"""
HỆ THỐNG ĐIỂM DANH CANTEEN SEAZEN - SERVER RUNNER & KIOSK BACKEND
Tích hợp Microsoft 365 SharePoint Online (List: Canteen_Orders & UserLogin)
Xác thực bắt buộc tài khoản Microsoft 365 (MSAL Bearer Token, AppData DPAPI Storage)
Đặc quyền VIP nhận diện duy nhất cho Tổng Giám Đốc ĐỖ VĂN TOÀN (Mã NV: 1, 01, 001, 0001, 00001)
"""

import os
import sys
import json
import datetime
import time
import webbrowser
import requests
import msal
import socket
import logging
import threading
import io
import secrets
import ctypes
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from flask import Flask, jsonify, request, send_file, send_from_directory

# Cấu hình stdout / stderr hiển thị tiếng Việt UTF-8 chuẩn trên Windows Console
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

# Tắt bớt log HTTP mặc định của Werkzeug để màn hình console sạch sẽ
log = logging.getLogger('werkzeug')
log.setLevel(logging.ERROR)

app = Flask(__name__)

# =============================================================================
# CẤU HÌNH THƯ MỤC THỰC THI & LƯU TRỮ TOKEN ẨN TẠI APPDATA
# =============================================================================
IS_FROZEN = getattr(sys, 'frozen', False)
if IS_FROZEN:
    EXE_DIR = os.path.dirname(sys.executable)
    BUNDLE_DIR = getattr(sys, '_MEIPASS', EXE_DIR)
else:
    EXE_DIR = os.path.dirname(os.path.abspath(__file__))
    BUNDLE_DIR = EXE_DIR

# Thư mục lưu trữ dữ liệu an toàn & ẩn trong AppData máy cục bộ
APPDATA_BASE = os.environ.get("LOCALAPPDATA") or os.environ.get("APPDATA") or os.path.expanduser("~")
DATA_DIR = os.path.join(APPDATA_BASE, "SeaZen_Canteen_Kiosk")
os.makedirs(DATA_DIR, exist_ok=True)
if os.name == 'nt':
    try:
        ctypes.windll.kernel32.SetFileAttributesW(DATA_DIR, 0x02)  # FILE_ATTRIBUTE_HIDDEN
    except Exception:
        pass

TOKEN_CACHE_FILE = os.path.join(DATA_DIR, "token_cache.bin")

def resolve_bundled_file(filename, fallback_dir=None):
    """Tìm file tài nguyên ưu tiên trong BUNDLE_DIR (nhúng bên trong EXE), sau đó mới tới EXE_DIR"""
    candidate_paths = [
        os.path.join(BUNDLE_DIR, filename),
        os.path.join(EXE_DIR, filename),
        os.path.join(EXE_DIR, "OneSEAZEN", filename)
    ]
    if fallback_dir:
        candidate_paths.append(os.path.join(fallback_dir, filename))
    for p in candidate_paths:
        if os.path.exists(p):
            return p
    return os.path.join(BUNDLE_DIR, filename)

HTML_FILE = resolve_bundled_file("Canteen_Kiosk_Production.html")
MP3_DIR = os.path.join(BUNDLE_DIR, "mp3") if os.path.isdir(os.path.join(BUNDLE_DIR, "mp3")) else os.path.join(EXE_DIR, "mp3")
LOGO_FILE = resolve_bundled_file("seazen_logo.png")
ICON_FILE = resolve_bundled_file("canteen_icon.ico")

# =============================================================================
# CẤU HÌNH MICROSOFT ENTRA ID & SHAREPOINT SITE ONESEAZEN
# =============================================================================
TENANT_ID = os.environ.get("MS_TENANT_ID", "29101901-3778-4339-827c-7d2b50914db0")
# App chính thức SEAZENHR-PORTAL theo Quy chuẩn no-cookies-auth.md
CLIENT_ID_PRIMARY = os.environ.get("MS_CLIENT_ID", "0390280b-12b8-49c4-adbf-a9c2b8bccfb5")
CLIENT_ID = CLIENT_ID_PRIMARY
CLIENT_ID_FALLBACKS = [
    "0390280b-12b8-49c4-adbf-a9c2b8bccfb5",  # SEAZENHR-PORTAL (Chính thức)
    "04b07795-8ddb-461a-bbee-02f9e1bf7b46",  # Azure CLI (Tương thích ngược token cũ trên máy người dùng)
    "67ab195e-f6ff-40ff-9169-73035faaf572",  # PnP-SeaZen
]

AUTHORITY = f"https://login.microsoftonline.com/{TENANT_ID}"
SCOPES = ["https://graph.microsoft.com/.default"]
SP_SCOPES = ["https://seazenvn.sharepoint.com/.default"]
SITE_ID = "seazenvn.sharepoint.com,d47e794f-870a-4d63-8173-4d0721d5d77d,69313196-298d-4be2-af46-f7893c08cccf"
LIST_NAME = "Canteen_Orders"
USER_LOGIN_LIST_NAME = "UserLogin"

VN_TZ = datetime.timezone(datetime.timedelta(hours=7))

ACTIVE_SESSIONS = {}
SESSION_TTL_SECONDS = 8 * 60 * 60
CHECKIN_LOCK = threading.Lock()
CACHE_LOCK = threading.RLock()
CACHED_USER_LOGINS = []
LAST_USER_LOGIN_FETCH = 0
USER_LOGIN_CACHE_TTL = 60

CONNECTION_DIAGNOSTICS = {
    'token_file_found': os.path.exists(TOKEN_CACHE_FILE),
    'token_file_path': TOKEN_CACHE_FILE,
    'last_error': '',
    'user_account': '',
    'user_name': '',
    'last_sync_time': ''
}

LOCAL_CHECKIN_OVERLAY = {}
LAST_FETCH_TIME = 0
CACHED_ORDERS_TODAY = {}
CACHED_CLAIMED_LIST = []
CACHED_ALL_TODAY_LIST = []
CACHE_TTL_SECONDS = 3
PENDING_DEVICE_FLOW = None

# =============================================================================
# MÃ HÓA WINDOWS DPAPI BẢO VỆ TOKEN TRÊN MÁY CỤC BỘ
# =============================================================================
def dpapi_encrypt(data_bytes: bytes) -> bytes:
    """Mã hóa token bằng Windows DPAPI (CryptProtectData) gắn với tài khoản người dùng Windows"""
    if not data_bytes or os.name != 'nt':
        return data_bytes
    try:
        from ctypes import wintypes
        class DATA_BLOB(ctypes.Structure):
            _fields_ = [('cbData', wintypes.DWORD), ('pbData', ctypes.POINTER(ctypes.c_byte))]

        data_in = DATA_BLOB(
            len(data_bytes),
            ctypes.cast(ctypes.create_string_buffer(data_bytes), ctypes.POINTER(ctypes.c_byte))
        )
        data_out = DATA_BLOB()
        if ctypes.windll.crypt32.CryptProtectData(
            ctypes.byref(data_in), "SeaZen Canteen Token Cache", None, None, None, 0, ctypes.byref(data_out)
        ):
            encrypted = ctypes.string_at(data_out.pbData, data_out.cbData)
            ctypes.windll.kernel32.LocalFree(data_out.pbData)
            return b"DPAPI:" + encrypted
    except Exception as e:
        print(f"[DPAPI] Lỗi mã hóa: {e}")
    return data_bytes


def dpapi_decrypt(data_bytes: bytes) -> bytes:
    """Giải mã token bằng Windows DPAPI (CryptUnprotectData)"""
    if not data_bytes or os.name != 'nt' or not data_bytes.startswith(b"DPAPI:"):
        return data_bytes
    try:
        from ctypes import wintypes
        class DATA_BLOB(ctypes.Structure):
            _fields_ = [('cbData', wintypes.DWORD), ('pbData', ctypes.POINTER(ctypes.c_byte))]

        raw_enc = data_bytes[6:]
        blob_in = DATA_BLOB(
            len(raw_enc),
            ctypes.cast(ctypes.create_string_buffer(raw_enc), ctypes.POINTER(ctypes.c_byte))
        )
        blob_out = DATA_BLOB()
        if ctypes.windll.crypt32.CryptUnprotectData(
            ctypes.byref(blob_in), None, None, None, None, 0, ctypes.byref(blob_out)
        ):
            decrypted = ctypes.string_at(blob_out.pbData, blob_out.cbData)
            ctypes.windll.kernel32.LocalFree(blob_out.pbData)
            return decrypted
    except Exception as e:
        print(f"[DPAPI] Lỗi giải mã: {e}")
    return data_bytes


# =============================================================================
# QUẢN LÝ VÒNG ĐỜI MSAL TOKEN CACHE
# =============================================================================
def init_msal_cache():
    """Nạp cache token từ AppData, hỗ trợ tự động di trú từ bản cũ nếu cần"""
    cache = msal.SerializableTokenCache()
    if os.path.exists(TOKEN_CACHE_FILE):
        try:
            with open(TOKEN_CACHE_FILE, "rb") as f:
                raw_bytes = f.read()
            decrypted = dpapi_decrypt(raw_bytes)
            cache.deserialize(decrypted.decode("utf-8", errors="replace"))
            return cache
        except Exception as e:
            print(f"[MSAL] Lỗi nạp token từ AppData: {e}")

    # Fallback 1: Đồng bộ từ AppData của SeaZen HR Portal nếu đã đăng nhập trên máy này
    portal_cache = os.path.join(APPDATA_BASE, "SeaZen_HR_Portal", "data", ".cache", "token_cache.bin")
    if os.path.exists(portal_cache):
        try:
            with open(portal_cache, "rb") as f:
                raw_bytes = f.read()
            decrypted = dpapi_decrypt(raw_bytes)
            cache.deserialize(decrypted.decode("utf-8", errors="replace"))
            save_token_cache(cache)
            print("[MSAL] Đồng bộ phiên đăng nhập từ SeaZen HR Portal [OK]")
            return cache
        except Exception:
            pass

    # Fallback 2: token_cache.bin cũ cạnh file .exe
    legacy_token = os.path.join(EXE_DIR, "token_cache.bin")
    if os.path.exists(legacy_token):
        try:
            with open(legacy_token, "rb") as f:
                raw_b = f.read()
            dec = dpapi_decrypt(raw_b)
            cache.deserialize(dec.decode("utf-8", errors="replace"))
            save_token_cache(cache)
            print("[MSAL] Đã di trú token từ thư mục gốc sang AppData [OK]")
            return cache
        except Exception:
            pass

    return cache


MSAL_CACHE = init_msal_cache()


def save_token_cache(cache=None):
    """Lưu token cache đã mã hóa DPAPI vào thư mục ẩn AppData"""
    with CACHE_LOCK:
        target = cache or MSAL_CACHE
        try:
            os.makedirs(os.path.dirname(TOKEN_CACHE_FILE), exist_ok=True)
            content = target.serialize().encode("utf-8")
            enc = dpapi_encrypt(content)
            with open(TOKEN_CACHE_FILE, "wb") as f:
                f.write(enc)
            CONNECTION_DIAGNOSTICS['token_file_found'] = True
        except Exception as e:
            print(f"[MSAL] Không thể lưu token_cache.bin: {e}")


def get_current_ms_account():
    """Lấy thông tin tài khoản Microsoft đang hoạt động trong token cache"""
    with CACHE_LOCK:
        for cid in CLIENT_ID_FALLBACKS:
            app_client = msal.PublicClientApplication(
                cid,
                authority=AUTHORITY,
                token_cache=MSAL_CACHE
            )
            accounts = app_client.get_accounts()
            if accounts:
                acc = accounts[0]
                u = acc.get("username", "")
                name = acc.get("name", "") or u
                CONNECTION_DIAGNOSTICS['user_account'] = u
                CONNECTION_DIAGNOSTICS['user_name'] = name
                return {"username": u, "name": name}
    return None


def get_access_token():
    """
    Lấy Microsoft Bearer Token qua acquire_token_silent ngầm.
    Nếu chưa đăng nhập hoặc token hết hạn hoàn toàn, trả về None để ép người dùng đăng nhập.
    """
    with CACHE_LOCK:
        for cid in CLIENT_ID_FALLBACKS:
            app_client = msal.PublicClientApplication(
                cid,
                authority=AUTHORITY,
                token_cache=MSAL_CACHE
            )
            accounts = app_client.get_accounts()
            if not accounts:
                continue

            acc = accounts[0]
            for attempt in range(2):
                try:
                    res = app_client.acquire_token_silent(SCOPES, account=acc)
                except Exception:
                    res = None
                if res and "access_token" in res:
                    if MSAL_CACHE.has_state_changed:
                        save_token_cache(MSAL_CACHE)
                    CONNECTION_DIAGNOSTICS['token_file_found'] = True
                    CONNECTION_DIAGNOSTICS['user_account'] = acc.get('username', '')
                    CONNECTION_DIAGNOSTICS['last_error'] = ''
                    return res["access_token"]
                if attempt == 0:
                    time.sleep(0.5)

        CONNECTION_DIAGNOSTICS['token_file_found'] = os.path.exists(TOKEN_CACHE_FILE)
        acc = get_current_ms_account()
        if not acc:
            CONNECTION_DIAGNOSTICS['last_error'] = "REQUIRE_MS_LOGIN"
        else:
            CONNECTION_DIAGNOSTICS['last_error'] = "TOKEN_EXPIRED"
        return None


def clear_token_cache():
    """Xóa tài khoản và token khỏi bộ nhớ và AppData khi đăng xuất"""
    with CACHE_LOCK:
        for cid in CLIENT_ID_FALLBACKS:
            try:
                app_client = msal.PublicClientApplication(cid, authority=AUTHORITY, token_cache=MSAL_CACHE)
                for a in app_client.get_accounts():
                    app_client.remove_account(a)
            except Exception:
                pass
        try:
            if os.path.exists(TOKEN_CACHE_FILE):
                os.remove(TOKEN_CACHE_FILE)
        except Exception:
            pass
        CONNECTION_DIAGNOSTICS['token_file_found'] = False
        CONNECTION_DIAGNOSTICS['user_account'] = ''
        CONNECTION_DIAGNOSTICS['user_name'] = ''
        CONNECTION_DIAGNOSTICS['last_error'] = 'REQUIRE_MS_LOGIN'


# =============================================================================
# CHUẨN HÓA MÃ SỐ NHÂN VIÊN (ĐỖ VĂN TOÀN: 1, 01, 001, 0001, 00001)
# =============================================================================
def normalize_emp_id(raw_val):
    """
    Chuẩn hóa mã nhân viên loại bỏ các số 0 ở đầu nếu là chuỗi số:
    Ví dụ: '1', '01', '001', '0001', '00001', '000001' đều quy về '1'.
    Hỗ trợ cả chuỗi mã vạch / QR mở rộng dạng '00001|Đỗ Văn Toàn|Tổng Giám Đốc'.
    """
    if raw_val is None:
        return ""
    s = str(raw_val).strip()
    if '|' in s:
        s = s.split('|', 1)[0].strip()
    if s.isdigit():
        return str(int(s))
    return s


def get_lan_ip():
    """Lấy địa chỉ IP nội bộ của máy trạm trong mạng LAN / Wi-Fi"""
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.settimeout(0.5)
        s.connect(('8.8.8.8', 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        try:
            return socket.gethostbyname(socket.gethostname())
        except Exception:
            return "127.0.0.1"


def get_today_vn_str():
    """Lấy ngày hôm nay theo múi giờ Việt Nam định dạng YYYY-MM-DD"""
    return datetime.datetime.now(VN_TZ).strftime('%Y-%m-%d')


def normalize_sp_date_to_vn(raw_date_str):
    """Chuẩn hóa ngày từ SharePoint sang định dạng YYYY-MM-DD theo giờ VN"""
    if not raw_date_str:
        return ""
    str_val = str(raw_date_str).strip()
    if len(str_val) == 10 and str_val[4] == '-' and str_val[7] == '-':
        return str_val
    try:
        clean = str_val.replace("Z", "+00:00")
        dt = datetime.datetime.fromisoformat(clean)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=datetime.timezone.utc)
        dt_vn = dt.astimezone(VN_TZ)
        return dt_vn.strftime('%Y-%m-%d')
    except Exception:
        return str(raw_date_str)[:10]


# =============================================================================
# ĐỒNG BỘ DỮ LIỆU TỪ SHAREPOINT ONLINE (CANTEEN_ORDERS)
# =============================================================================
def fetch_sharepoint_data(force=False):
    """Đồng bộ danh sách suất ăn hôm nay từ SharePoint List Canteen_Orders"""
    global LAST_FETCH_TIME, CACHED_ORDERS_TODAY, CACHED_CLAIMED_LIST, CACHED_ALL_TODAY_LIST
    now_ts = time.time()
    today_vn = get_today_vn_str()

    if not force and CACHED_ORDERS_TODAY and (now_ts - LAST_FETCH_TIME < CACHE_TTL_SECONDS):
        for ma_nv, overlay in list(LOCAL_CHECKIN_OVERLAY.items()):
            if overlay.get('date') == today_vn:
                norm_ma = normalize_emp_id(ma_nv)
                for key in (ma_nv, norm_ma):
                    if key in CACHED_ORDERS_TODAY:
                        CACHED_ORDERS_TODAY[key]['status'] = 'Đã nhận cơm'
                        if not CACHED_ORDERS_TODAY[key].get('claimedAt'):
                            CACHED_ORDERS_TODAY[key]['claimedAt'] = overlay.get('claimedAt', '')
        return CACHED_ORDERS_TODAY, CACHED_CLAIMED_LIST, CACHED_ALL_TODAY_LIST

    token = get_access_token()
    if not token:
        return CACHED_ORDERS_TODAY or {}, CACHED_CLAIMED_LIST or [], CACHED_ALL_TODAY_LIST or []

    headers = {
        'Authorization': f'Bearer {token}',
        'Accept': 'application/json'
    }
    url = f"https://graph.microsoft.com/v1.0/sites/{SITE_ID}/lists/{LIST_NAME}/items?$expand=fields&$top=999"

    try:
        r = requests.get(url, headers=headers, timeout=10)
        if r.status_code == 200:
            CONNECTION_DIAGNOSTICS['last_error'] = ''
            CONNECTION_DIAGNOSTICS['last_sync_time'] = datetime.datetime.now(VN_TZ).strftime('%H:%M:%S')
            data = r.json()
            items = data.get('value', [])
            next_link = data.get('@odata.nextLink')
            while next_link:
                try:
                    r_next = requests.get(next_link, headers=headers, timeout=10)
                    if r_next.status_code == 200:
                        d_next = r_next.json()
                        items.extend(d_next.get('value', []))
                        next_link = d_next.get('@odata.nextLink')
                    else:
                        break
                except Exception:
                    break

            orders_today = {}
            claimed_list = []
            all_today_list = []

            for item in items:
                f = item.get('fields', {})
                ma_nv_raw = str(f.get('Title', '')).strip()
                norm_id = normalize_emp_id(ma_nv_raw)
                ho_ten = str(f.get('HoTen', '')).strip()
                chuc_vu = str(f.get('ChucVu', '')).strip()
                loai_mon = str(f.get('LoaiMon') or 'Món mặn').strip()
                try:
                    so_luong = int(f.get('SoLuongSuat') or 1)
                except Exception:
                    so_luong = 1
                trang_thai = str(f.get('TrangThai') or 'Đã đăng ký').strip()
                ngay_an_raw = str(f.get('NgayAn', '')).strip()
                gio_nhan = str(f.get('GioNhan', '')).strip()
                ghi_chu = str(f.get('GhiChu', '')).strip()

                order_date_vn = normalize_sp_date_to_vn(ngay_an_raw)
                try:
                    display_ngay_an = datetime.datetime.strptime(order_date_vn, '%Y-%m-%d').strftime('%d/%m/%Y')
                except Exception:
                    display_ngay_an = order_date_vn

                # Kiểm tra overlay cục bộ
                overlay = LOCAL_CHECKIN_OVERLAY.get(ma_nv_raw) or LOCAL_CHECKIN_OVERLAY.get(norm_id)
                if overlay and overlay.get('date') == today_vn:
                    trang_thai = 'Đã nhận cơm'
                    if not gio_nhan:
                        gio_nhan = overlay.get('claimedAt', '')

                if order_date_vn == today_vn and trang_thai != 'Đã hủy':
                    order_obj = {
                        'itemId': item.get('id'),
                        'empId': ma_nv_raw,
                        'normEmpId': norm_id,
                        'fullName': ho_ten,
                        'role': chuc_vu,
                        'mealType': loai_mon,
                        'quantity': so_luong,
                        'status': trang_thai,
                        'claimedAt': gio_nhan or '',
                        'orderDate': order_date_vn,
                        'mealDate': display_ngay_an,
                        'note': ghi_chu
                    }
                    # Đánh chỉ mục cả mã nguyên bản và mã đã chuẩn hóa để tra cứu tức thì
                    orders_today[ma_nv_raw] = order_obj
                    if norm_id:
                        orders_today[norm_id] = order_obj
                    all_today_list.append(order_obj)

                    if trang_thai == 'Đã nhận cơm':
                        claimed_list.append({
                            'name': ho_ten,
                            'empId': ma_nv_raw,
                            'normEmpId': norm_id,
                            'meal': f"{so_luong} {loai_mon}",
                            'time': gio_nhan or '11:30',
                            'note': ghi_chu
                        })

            claimed_list.sort(key=lambda x: str(x.get('time', '')), reverse=True)

            CACHED_ORDERS_TODAY = orders_today
            CACHED_CLAIMED_LIST = claimed_list
            CACHED_ALL_TODAY_LIST = all_today_list
            LAST_FETCH_TIME = now_ts
            return orders_today, claimed_list, all_today_list
        else:
            CONNECTION_DIAGNOSTICS['last_error'] = f"HTTP_{r.status_code}"
    except Exception as e:
        CONNECTION_DIAGNOSTICS['last_error'] = f"LOI_MANG: {str(e)[:40]}"

    return CACHED_ORDERS_TODAY or {}, CACHED_CLAIMED_LIST or [], CACHED_ALL_TODAY_LIST or []


def save_checkin_to_sharepoint(item_id, claimed_at, meal_date, vip=False):
    """Ghi nhận trên SharePoint trước khi báo check-in thành công."""
    token = get_access_token()
    if not token:
        return None, 'Không thể xác thực SharePoint. Yêu cầu đăng nhập tài khoản Microsoft 365.'

    headers = {'Authorization': f'Bearer {token}', 'Content-Type': 'application/json'}
    base_url = f"https://graph.microsoft.com/v1.0/sites/{SITE_ID}/lists/{LIST_NAME}/items"
    try:
        if vip and not item_id:
            # Tra cứu đơn của Tổng Giám Đốc nếu chưa có itemId
            url = f"{base_url}?$expand=fields&$top=999"
            while url:
                lookup = requests.get(url, headers=headers, timeout=10)
                if lookup.status_code != 200:
                    return None, f'Không kiểm tra được đơn trên SharePoint (HTTP {lookup.status_code}).'
                data = lookup.json()
                for item in data.get('value', []):
                    fields = item.get('fields', {})
                    t_val = str(fields.get('Title', '')).strip()
                    if normalize_emp_id(t_val) == '1' and normalize_sp_date_to_vn(fields.get('NgayAn')) == meal_date:
                        item_id = item.get('id')
                        break
                if item_id:
                    break
                url = data.get('@odata.nextLink')

        if item_id:
            response = requests.patch(
                f"{base_url}/{item_id}/fields", headers=headers,
                json={'TrangThai': 'Đã nhận cơm', 'GioNhan': claimed_at}, timeout=8
            )
            if response.status_code in (200, 204):
                return item_id, None
        elif vip:
            # Tự động tạo suất ăn danh dự cho Tổng Giám Đốc nếu chưa đăng ký trước
            response = requests.post(base_url, headers=headers, json={'fields': {
                'Title': '1', 'HoTen': 'ĐỖ VĂN TOÀN', 'ChucVu': 'Tổng Giám Đốc', 'TrangThai': 'Đã nhận cơm',
                'NgayAn': f'{meal_date}T00:00:00Z', 'GioNhan': claimed_at,
                'LoaiMon': 'Món mặn', 'SoLuongSuat': 1,
                'GhiChu': 'Suất ăn Tổng Giám Đốc (Quét trực tiếp tại Kiosk)'
            }}, timeout=8)
            if response.status_code in (200, 201):
                return response.json().get('id'), None
        else:
            return None, 'Đơn ăn thiếu mã SharePoint. Vui lòng làm mới danh sách.'
        return None, f'Ghi nhận trên SharePoint thất bại (HTTP {response.status_code}).'
    except requests.RequestException as exc:
        print(f'[-] Lỗi đồng bộ SharePoint: {exc}')
        return None, 'Không kết nối được SharePoint. Vui lòng thử lại.'


# =============================================================================
# ĐIỀU PHỐI ĐIỂM DANH (CHECK-IN ROUTE)
# =============================================================================
@app.route('/api/checkin', methods=['POST'])
def checkin():
    with CHECKIN_LOCK:
        return _checkin_locked()


def _checkin_locked():
    # Kiểm tra trạng thái xác thực tài khoản Microsoft
    token = get_access_token()
    if not token:
        return jsonify({
            'status': 'REQUIRE_MS_LOGIN',
            'message': 'Hệ thống chưa đăng nhập tài khoản Microsoft 365. Vui lòng đăng nhập để sử dụng.',
            'authenticated': False
        }), 401

    body = request.get_json(force=True, silent=True) or {}
    raw_ma_nv = str(body.get('maNV', '')).strip()
    raw_scan = str(body.get('rawScan', '')).strip()
    now_vn = datetime.datetime.now(VN_TZ)
    gio_nhan = now_vn.strftime('%H:%M:%S')

    candidate = raw_ma_nv or raw_scan
    candidate_lower = candidate.lower()

    # Kiểm tra mã độc hoặc URL không hợp lệ
    if (candidate_lower.startswith('http://') or candidate_lower.startswith('https://')
        or candidate_lower.startswith('www.') or candidate_lower.startswith('wifi:')
        or candidate_lower.startswith('000201') or len(candidate) > 100
        or any(c in '<>{}\\/?&=$%#~^";' for c in candidate if c != '|')):
        print(f"[-] [{gio_nhan}] TỪ CHỐI: Mã quét sai định dạng thẻ: {candidate[:30]}...")
        return jsonify({
            'status': 'INVALID_FORMAT',
            'message': 'Mã quét sai định dạng thẻ nhân viên',
            'empId': candidate[:30]
        }), 400

    scanned_id_part = candidate.split('|', 1)[0].strip()
    norm_scanned_id = normalize_emp_id(scanned_id_part)
    norm_raw_ma = normalize_emp_id(raw_ma_nv) if raw_ma_nv else norm_scanned_id

    # Xác thực tính nhất quán giữa rawScan và maNV
    if raw_scan and raw_ma_nv and norm_scanned_id != norm_raw_ma:
        return jsonify({'status': 'INVALID_FORMAT', 'message': 'Mã quét và mã nhân viên không khớp'}), 400

    if not norm_scanned_id or not scanned_id_part.isdigit() or not 1 <= len(scanned_id_part) <= 10:
        return jsonify({'status': 'INVALID_FORMAT', 'message': 'Mã nhân viên phải là chuỗi số từ 1 đến 10 chữ số'}), 400

    # =========================================================================
    # ĐẶC QUYỀN VIP CHO TỔNG GIÁM ĐỐC: ĐỖ VĂN TOÀN
    # BẤT KỲ MÃ SỐ: 1, 01, 001, 0001, 00001 ĐỀU LÀ 1 NGƯỜI DUY NHẤT
    # =========================================================================
    is_tgd = (norm_scanned_id == '1')

    if is_tgd:
        ma_nv = "1"
    else:
        ma_nv = norm_scanned_id

    orders_today = CACHED_ORDERS_TODAY
    if not orders_today:
        orders_today, _, _ = fetch_sharepoint_data()

    today_vn = get_today_vn_str()

    if is_tgd:
        print(f"[+] [{gio_nhan}] >>> NHẬN DIỆN TỔNG GIÁM ĐỐC ĐỖ VĂN TOÀN (MÃ NV: {scanned_id_part} -> 1) - DUYỆT VIP! <<<")

        # Tìm đơn ăn của anh Toàn trong cache
        matched_order = None
        if orders_today and ("1" in orders_today):
            matched_order = orders_today["1"]
        else:
            for o_item in CACHED_ALL_TODAY_LIST:
                if normalize_emp_id(o_item.get('empId', '')) == "1":
                    matched_order = o_item
                    break

        if matched_order and matched_order.get('status') == 'Đã nhận cơm':
            return jsonify({'status': 'ALREADY_CLAIMED', 'message': 'Đã nhận cơm trước đó', 'order': matched_order})

        item_id, sync_error = save_checkin_to_sharepoint(
            matched_order.get('itemId') if matched_order else None, gio_nhan, today_vn, vip=True
        )
        if sync_error:
            return jsonify({'status': 'SYNC_ERROR', 'message': sync_error, 'spSynced': False}), 503

        if matched_order:
            matched_order['itemId'] = item_id
            matched_order['status'] = 'Đã nhận cơm'
            matched_order['claimedAt'] = gio_nhan
            final_order = matched_order
        else:
            final_order = {
                'itemId': item_id,
                'empId': '1',
                'normEmpId': '1',
                'fullName': 'ĐỖ VĂN TOÀN',
                'role': 'Tổng Giám Đốc',
                'mealType': 'Món mặn',
                'quantity': 1,
                'status': 'Đã nhận cơm',
                'claimedAt': gio_nhan,
                'orderDate': today_vn,
                'mealDate': datetime.datetime.strptime(today_vn, '%Y-%m-%d').strftime('%d/%m/%Y'),
                'note': 'Suất ăn Tổng Giám Đốc (Quét trực tiếp tại Kiosk)'
            }

        orders_today["1"] = final_order
        CACHED_ALL_TODAY_LIST.append(final_order)

        # Cập nhật cache cục bộ
        LOCAL_CHECKIN_OVERLAY["1"] = {'date': today_vn, 'status': 'Đã nhận cơm', 'claimedAt': gio_nhan}
        if scanned_id_part != "1":
            LOCAL_CHECKIN_OVERLAY[scanned_id_part] = {'date': today_vn, 'status': 'Đã nhận cơm', 'claimedAt': gio_nhan}

        # Cập nhật feed nhận cơm gần đây
        global CACHED_CLAIMED_LIST
        CACHED_CLAIMED_LIST = [c for c in CACHED_CLAIMED_LIST if normalize_emp_id(c.get('empId', '')) != "1"]
        CACHED_CLAIMED_LIST.insert(0, {
            'name': 'ĐỖ VĂN TOÀN',
            'empId': '1',
            'meal': '1 Món mặn',
            'time': gio_nhan,
            'note': 'Tổng Giám Đốc'
        })

        return jsonify({
            'status': 'SUCCESS',
            'message': 'Kính mời Tổng Giám Đốc nhận cơm, chúc anh ngon miệng!',
            'order': final_order,
            'spSynced': True
        })

    # =========================================================================
    # LOGIC CHECKIN NHÂN VIÊN THÔNG THƯỜNG
    # =========================================================================
    matched_order = orders_today.get(norm_scanned_id) or orders_today.get(scanned_id_part)
    if not matched_order:
        print(f"[-] [{gio_nhan}] TỪ CHỐI: Mã NV {norm_scanned_id} chưa đăng ký cơm hôm nay.")
        return jsonify({
            'status': 'NOT_REGISTERED',
            'message': 'Chưa đăng ký cơm hôm nay',
            'empId': norm_scanned_id
        })

    # Chặn quét trước 11h30 (Chỉ áp dụng với nhân viên thông thường)
    is_before_1130 = (now_vn.hour < 11 or (now_vn.hour == 11 and now_vn.minute < 30))
    if is_before_1130:
        print(f"[!] [{gio_nhan}] TỪ CHỐI: Quét trước 11h30 (Hiện tại: {gio_nhan}). Mã NV: {norm_scanned_id}")
        return jsonify({
            'status': 'BEFORE_TIME',
            'message': 'Chưa đến giờ nhận cơm! Khung giờ phục vụ bắt đầu từ 11:30 hôm nay',
            'currentTime': gio_nhan,
            'order': matched_order,
            'empId': norm_scanned_id
        })

    # Đã nhận cơm trước đó
    if matched_order.get('status') == 'Đã nhận cơm':
        print(f"[!] [{gio_nhan}] CẢNH BÁO: Mã NV {norm_scanned_id} ({matched_order.get('fullName')}) đã nhận cơm trước đó lúc {matched_order.get('claimedAt')}.")
        return jsonify({
            'status': 'ALREADY_CLAIMED',
            'message': 'Đã nhận cơm trước đó',
            'order': matched_order
        })

    # Ghi nhận trên SharePoint trước khi xác nhận
    _, sync_error = save_checkin_to_sharepoint(matched_order.get('itemId'), gio_nhan, today_vn)
    if sync_error:
        return jsonify({'status': 'SYNC_ERROR', 'message': sync_error, 'spSynced': False}), 503

    # Cập nhật trạng thái
    LOCAL_CHECKIN_OVERLAY[norm_scanned_id] = {
        'date': today_vn,
        'status': 'Đã nhận cơm',
        'claimedAt': gio_nhan
    }
    if scanned_id_part != norm_scanned_id:
        LOCAL_CHECKIN_OVERLAY[scanned_id_part] = LOCAL_CHECKIN_OVERLAY[norm_scanned_id]

    matched_order['status'] = 'Đã nhận cơm'
    matched_order['claimedAt'] = gio_nhan

    for item in CACHED_ALL_TODAY_LIST:
        if normalize_emp_id(item.get('empId')) == norm_scanned_id:
            item['status'] = 'Đã nhận cơm'
            item['claimedAt'] = gio_nhan
            break

    CACHED_CLAIMED_LIST = [c for c in CACHED_CLAIMED_LIST if normalize_emp_id(c.get('empId')) != norm_scanned_id]
    CACHED_CLAIMED_LIST.insert(0, {
        'name': matched_order.get('fullName'),
        'empId': matched_order.get('empId'),
        'meal': f"{matched_order.get('quantity')} {matched_order.get('mealType')}",
        'time': gio_nhan,
        'note': matched_order.get('note', '')
    })

    print(f"[+] [{gio_nhan}] CHECK-IN THÀNH CÔNG: Mã NV {norm_scanned_id} | {matched_order.get('fullName')} | {matched_order.get('quantity')} suất ({matched_order.get('mealType')})")

    return jsonify({
        'status': 'SUCCESS',
        'message': 'Check-in thành công',
        'order': matched_order,
        'spSynced': True
    })


# =============================================================================
# ENDPOINTS XÁC THỰC MICROSOFT 365 (BẮT BUỘC)
# =============================================================================
@app.route('/api/auth/ms-status', methods=['GET'])
def api_ms_status():
    """Kiểm tra trạng thái xác thực Microsoft 365"""
    token = get_access_token()
    acc = get_current_ms_account()
    return jsonify({
        'authenticated': bool(token),
        'account': acc.get('username', '') if acc else '',
        'name': acc.get('name', '') if acc else '',
        'tokenFound': CONNECTION_DIAGNOSTICS['token_file_found'],
        'lastError': CONNECTION_DIAGNOSTICS['last_error'],
        'hasPendingDeviceFlow': bool(PENDING_DEVICE_FLOW)
    })


@app.route('/api/auth/ms-login-interactive', methods=['POST'])
def api_ms_login_interactive():
    """Đăng nhập tương tác bằng trình duyệt hệ thống qua MSAL OAuth2 (Redirect localhost)"""
    try:
        app_client = msal.PublicClientApplication(
            CLIENT_ID_PRIMARY,
            authority=AUTHORITY,
            token_cache=MSAL_CACHE
        )
        res = app_client.acquire_token_interactive(
            scopes=SCOPES,
            prompt="select_account",
            timeout=120
        )
        if res and "access_token" in res:
            save_token_cache(MSAL_CACHE)
            acc = get_current_ms_account()
            fetch_sharepoint_data(force=True)
            print(f"[+] ĐĂNG NHẬP MICROSOFT 365 THÀNH CÔNG: {acc.get('username') if acc else ''}")
            return jsonify({
                'status': 'SUCCESS',
                'message': 'Đăng nhập Microsoft 365 thành công!',
                'account': acc
            })
        err_msg = res.get('error_description') or res.get('error') or 'Đăng nhập không thành công hoặc đã bị hủy.'
        return jsonify({'status': 'ERROR', 'message': err_msg}), 400
    except Exception as e:
        return jsonify({'status': 'ERROR', 'message': f'Lỗi xác thực: {str(e)}'}), 500


@app.route('/api/auth/ms-device-code', methods=['POST'])
def api_ms_device_code():
    """Khởi tạo Device Code Flow hiển thị mã xác nhận trên màn hình Kiosk"""
    global PENDING_DEVICE_FLOW
    try:
        app_client = msal.PublicClientApplication(
            CLIENT_ID_PRIMARY,
            authority=AUTHORITY,
            token_cache=MSAL_CACHE
        )
        flow = app_client.initiate_device_flow(scopes=SCOPES)
        if "user_code" not in flow:
            return jsonify({'status': 'ERROR', 'message': 'Không thể khởi tạo Device Flow'}), 500
        PENDING_DEVICE_FLOW = flow
        # Tự động mở trình duyệt nếu có thể
        try:
            webbrowser.open(flow.get('verification_uri'))
        except Exception:
            pass
        return jsonify({
            'status': 'SUCCESS',
            'userCode': flow.get('user_code'),
            'verificationUri': flow.get('verification_uri'),
            'message': flow.get('message', ''),
            'expiresIn': flow.get('expires_in', 900)
        })
    except Exception as e:
        return jsonify({'status': 'ERROR', 'message': str(e)}), 500


@app.route('/api/auth/ms-device-poll', methods=['POST'])
def api_ms_device_poll():
    """Kiểm tra người dùng đã hoàn tất xác nhận Device Code trên trình duyệt chưa"""
    global PENDING_DEVICE_FLOW
    if not PENDING_DEVICE_FLOW:
        return jsonify({'status': 'ERROR', 'message': 'Không có phiên xác thực mã thiết bị.'}), 400
    try:
        app_client = msal.PublicClientApplication(
            CLIENT_ID_PRIMARY,
            authority=AUTHORITY,
            token_cache=MSAL_CACHE
        )
        res = app_client.acquire_token_by_device_flow(PENDING_DEVICE_FLOW)
        if res and "access_token" in res:
            save_token_cache(MSAL_CACHE)
            PENDING_DEVICE_FLOW = None
            acc = get_current_ms_account()
            fetch_sharepoint_data(force=True)
            return jsonify({'status': 'SUCCESS', 'account': acc})
        err = res.get('error') or ''
        if err == 'authorization_pending':
            return jsonify({'status': 'PENDING', 'message': 'Đang chờ xác thực trên trình duyệt...'})
        PENDING_DEVICE_FLOW = None
        return jsonify({'status': 'ERROR', 'message': res.get('error_description', 'Xác thực thất bại')}), 400
    except Exception as e:
        return jsonify({'status': 'ERROR', 'message': str(e)}), 500


@app.route('/api/auth/ms-login-credentials', methods=['POST'])
def api_ms_login_credentials():
    """Đăng nhập trực tiếp bằng Email & Mật khẩu tài khoản Microsoft 365"""
    data = request.get_json(force=True, silent=True) or {}
    username = str(data.get('username', '')).strip()
    password = str(data.get('password', '')).strip()
    if not username or not password:
        return jsonify({'status': 'ERROR', 'message': 'Vui lòng nhập đầy đủ Email và Mật khẩu Microsoft!'}), 400

    try:
        app_client = msal.PublicClientApplication(
            CLIENT_ID_PRIMARY,
            authority=AUTHORITY,
            token_cache=MSAL_CACHE
        )
        res = app_client.acquire_token_by_username_password(username, password, scopes=SCOPES)
        if res and "access_token" in res:
            save_token_cache(MSAL_CACHE)
            acc = get_current_ms_account()
            fetch_sharepoint_data(force=True)
            return jsonify({'status': 'SUCCESS', 'account': acc})
        err_desc = res.get('error_description', '')
        if any(k in err_desc for k in ['AADSTS50076', 'AADSTS50079', 'MFA', 'multi-factor']):
            return jsonify({
                'status': 'MFA_REQUIRED',
                'message': 'Tài khoản yêu cầu xác thực 2 bước (MFA). Vui lòng sử dụng tính năng "Đăng nhập bằng Microsoft 365" trên trình duyệt.'
            }), 400
        return jsonify({'status': 'ERROR', 'message': err_desc or res.get('error', 'Đăng nhập không thành công')}), 400
    except Exception as e:
        return jsonify({'status': 'ERROR', 'message': str(e)}), 500


@app.route('/api/auth/ms-logout', methods=['POST'])
def api_ms_logout():
    """Đăng xuất tài khoản Microsoft 365 và xóa token ẩn trong AppData"""
    clear_token_cache()
    global CACHED_ORDERS_TODAY, CACHED_CLAIMED_LIST, CACHED_ALL_TODAY_LIST
    CACHED_ORDERS_TODAY = {}
    CACHED_CLAIMED_LIST = []
    CACHED_ALL_TODAY_LIST = []
    return jsonify({'status': 'SUCCESS', 'message': 'Đã đăng xuất tài khoản Microsoft 365'})


# =============================================================================
# STATIC ASSETS & GIAO DIỆN KIOSK
# =============================================================================
@app.route('/')
def index():
    if not os.path.exists(HTML_FILE):
        return f"<h3>Lỗi hệ thống: Không tìm thấy tệp giao diện Kiosk ({HTML_FILE})</h3>", 500
    return send_file(HTML_FILE)


@app.route('/mp3/<path:filename>')
def serve_mp3(filename):
    for target_dir in [MP3_DIR, os.path.join(BUNDLE_DIR, 'mp3'), os.path.join(EXE_DIR, 'mp3')]:
        p = os.path.join(target_dir, filename)
        if os.path.exists(p):
            return send_file(p)
    return jsonify({'error': 'Audio not found'}), 404


@app.route('/seazen_logo.png')
def serve_logo():
    for target in [LOGO_FILE, os.path.join(BUNDLE_DIR, 'seazen_logo.png'), os.path.join(EXE_DIR, 'seazen_logo.png')]:
        if os.path.exists(target):
            return send_file(target)
    return jsonify({'error': 'Logo not found'}), 404


@app.route('/canteen_icon.ico')
def serve_icon():
    for target in [ICON_FILE, os.path.join(BUNDLE_DIR, 'canteen_icon.ico'), os.path.join(EXE_DIR, 'canteen_icon.ico')]:
        if os.path.exists(target):
            return send_file(target)
    return jsonify({'error': 'Icon not found'}), 404


@app.route('/img/<path:filename>')
def serve_img(filename):
    for target_dir in [BUNDLE_DIR, EXE_DIR, os.path.join(BUNDLE_DIR, 'img'), os.path.join(EXE_DIR, 'img')]:
        p = os.path.join(target_dir, filename)
        if os.path.exists(p):
            return send_file(p)
    return jsonify({'error': 'Image not found'}), 404


@app.route('/api/orders/today', methods=['GET'])
def get_orders_today():
    force_refresh = request.args.get('force', '0') in ('1', 'true', 'True')
    if force_refresh:
        LOCAL_CHECKIN_OVERLAY.clear()

    token = get_access_token()
    if not token:
        return jsonify({
            'status': 'REQUIRE_MS_LOGIN',
            'authenticated': False,
            'message': 'Yêu cầu đăng nhập tài khoản Microsoft 365 để sử dụng hệ thống.',
            'date': get_today_vn_str(),
            'total': 0,
            'orders': {},
            'recentClaimed': [],
            'allOrders': []
        })

    orders, claimed_list, all_today = fetch_sharepoint_data(force=force_refresh)
    has_error = bool(CONNECTION_DIAGNOSTICS['last_error'])
    return jsonify({
        'status': 'ERROR' if has_error and not orders else 'SUCCESS',
        'authenticated': True,
        'errorMessage': CONNECTION_DIAGNOSTICS['last_error'],
        'account': CONNECTION_DIAGNOSTICS['user_account'],
        'userName': CONNECTION_DIAGNOSTICS['user_name'],
        'tokenFound': CONNECTION_DIAGNOSTICS['token_file_found'],
        'date': get_today_vn_str(),
        'total': len(all_today),
        'orders': orders,
        'recentClaimed': claimed_list,
        'allOrders': all_today
    })


# =============================================================================
# XUẤT BÁO CÁO EXCEL CHUẨN DOANH NGHIỆP
# =============================================================================
def get_authenticated_user():
    token = request.headers.get('X-Auth-Token', '')
    session = ACTIVE_SESSIONS.get(token) if token else None
    if not session:
        return None
    if time.time() - session['createdAt'] > SESSION_TTL_SECONDS:
        ACTIVE_SESSIONS.pop(token, None)
        return None
    return session['user']


@app.route('/api/export-excel', methods=['POST'])
def export_excel():
    # Cho phép xuất báo cáo nếu phiên Microsoft hợp lệ hoặc có phiên quản trị
    token = get_access_token()
    if not token and not get_authenticated_user():
        return jsonify({'status': 'UNAUTHORIZED', 'message': 'Vui lòng đăng nhập tài khoản Microsoft để xuất báo cáo Excel.'}), 401

    data = request.get_json(force=True, silent=True) or {}
    date_str = str(data.get('date') or get_today_vn_str()).strip()
    orders = data.get('orders')
    if orders is None:
        _, _, orders = fetch_sharepoint_data()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Diem_Danh_Suat_An"
    ws.views.sheetView[0].showGridLines = True

    NAVY_THEME = "1B263B"
    BORDER_COLOR = "CBD5E1"
    GREEN_FILL = "DCFCE7"
    GREEN_TEXT = "15803D"
    AMBER_FILL = "FEF3C7"
    AMBER_TEXT = "B45309"

    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )

    ws.row_dimensions[2].height = 42
    ws.merge_cells('A2:F2')
    title_cell = ws['A2']
    title_cell.value = "BÁO CÁO CHI TIẾT SUẤT ĂN CANTEEN SEAZEN"
    title_cell.font = Font(name='Segoe UI', size=15, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color=NAVY_THEME, end_color=NAVY_THEME, fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')

    ws.row_dimensions[4].height = 20
    c_date = ws['B4']
    c_date.value = f"Ngày báo cáo: {date_str}"
    c_date.font = Font(name='Segoe UI', size=11, bold=True, color='334155')

    total_count = len(orders)
    claimed_count = sum(1 for o in orders if ('nhận' in str(o.get('status', '')).lower() or 'claim' in str(o.get('status', '')).lower()))
    waiting_count = total_count - claimed_count

    c_summary = ws['D4']
    c_summary.value = f"Tổng: {total_count} suất | Đã nhận: {claimed_count} | Chưa nhận: {waiting_count}"
    c_summary.font = Font(name='Segoe UI', size=11, bold=True, color='1E293B')

    headers = ['STT', 'MÃ NHÂN VIÊN', 'HỌ VÀ TÊN', 'TRẠNG THÁI', 'GIỜ NHẬN', 'GHI CHÚ']
    ws.row_dimensions[7].height = 28
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=7, column=col_idx, value=h_text)
        cell.font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color=NAVY_THEME, end_color=NAVY_THEME, fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    current_row = 8
    for idx, item in enumerate(orders, start=1):
        ws.row_dimensions[current_row].height = 24
        row_bg = "F8FAFC" if idx % 2 == 0 else "FFFFFF"

        stt = idx
        emp_id = str(item.get('empId') or item.get('Title') or '')
        full_name = str(item.get('fullName') or item.get('name') or '')
        status = str(item.get('status') or 'Chưa nhận')
        claimed_at = str(item.get('claimedAt') or item.get('time') or '')
        note = str(item.get('note') or '')
        is_claimed = ('nhận' in status.lower() or 'claim' in status.lower())

        c1 = ws.cell(row=current_row, column=1, value=stt)
        c1.alignment = Alignment(horizontal='center', vertical='center')
        c1.font = Font(name='Segoe UI', size=10.5, color='64748B')

        c2 = ws.cell(row=current_row, column=2, value=emp_id)
        c2.alignment = Alignment(horizontal='center', vertical='center')
        c2.font = Font(name='Consolas', size=11, bold=True, color='1E3A8A')

        c3 = ws.cell(row=current_row, column=3, value=full_name)
        c3.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        c3.font = Font(name='Segoe UI', size=11, bold=True, color='0F172A')

        status_fill = GREEN_FILL if is_claimed else AMBER_FILL
        status_text_color = GREEN_TEXT if is_claimed else AMBER_TEXT
        status_label = "✓ Đã nhận cơm" if is_claimed else "⏳ Chưa nhận"

        c4 = ws.cell(row=current_row, column=4, value=status_label)
        c4.alignment = Alignment(horizontal='center', vertical='center')
        c4.font = Font(name='Segoe UI', size=10.5, bold=True, color=status_text_color)
        c4.fill = PatternFill(start_color=status_fill, end_color=status_fill, fill_type='solid')

        time_val = claimed_at if (is_claimed and claimed_at) else "-"
        c5 = ws.cell(row=current_row, column=5, value=time_val)
        c5.alignment = Alignment(horizontal='center', vertical='center')
        c5.font = Font(name='Consolas', size=10.5, color='334155' if time_val != '-' else '94A3B8')

        c6 = ws.cell(row=current_row, column=6, value=note or "-")
        c6.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        c6.font = Font(name='Segoe UI', size=10.5, color='475569')

        for col_i in range(1, 7):
            c = ws.cell(row=current_row, column=col_i)
            c.border = thin_border
            if col_i != 4 and row_bg != "FFFFFF":
                c.fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type='solid')

        current_row += 1

    widths = {'A': 8, 'B': 18, 'C': 34, 'D': 18, 'E': 16, 'F': 30}
    for col_letter, w in widths.items():
        ws.column_dimensions[col_letter].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    clean_date = date_str.replace('/', '_').replace(' ', '_').replace('-', '_')
    filename = f"Danh_Sach_Suat_An_{clean_date}.xlsx"
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )


def fetch_user_logins(force=False):
    """Lấy danh sách tài khoản phân quyền Kiosk từ SharePoint List UserLogin"""
    global CACHED_USER_LOGINS, LAST_USER_LOGIN_FETCH
    now_ts = time.time()
    if not force and CACHED_USER_LOGINS and (now_ts - LAST_USER_LOGIN_FETCH < USER_LOGIN_CACHE_TTL):
        return CACHED_USER_LOGINS

    token = get_access_token()
    if not token:
        return CACHED_USER_LOGINS or []

    headers = {'Authorization': f'Bearer {token}', 'Accept': 'application/json'}
    url = f"https://graph.microsoft.com/v1.0/sites/{SITE_ID}/lists/{USER_LOGIN_LIST_NAME}/items?$expand=fields&$top=999"
    try:
        r = requests.get(url, headers=headers, timeout=10)
        if r.status_code == 200:
            data = r.json()
            items = data.get('value', [])
            parsed_users = []
            for it in items:
                f = it.get('fields', {})
                u = str(f.get('Username') or f.get('Title') or '').strip()
                p = str(f.get('Password') or '').strip()
                roles = f.get('AdminRoles') or []
                if isinstance(roles, str):
                    roles = [roles]
                elif not isinstance(roles, list):
                    roles = []
                parsed_users.append({
                    'username': u,
                    'password': p,
                    'fullName': str(f.get('FullName') or f.get('HoTen') or u).strip(),
                    'department': str(f.get('Department') or f.get('PhongBan') or '').strip(),
                    'securityRole': str(f.get('SecurityRole') or '').strip(),
                    'adminRoles': roles
                })
            CACHED_USER_LOGINS = parsed_users
            LAST_USER_LOGIN_FETCH = now_ts
            return parsed_users
    except Exception as e:
        print("[-] Lỗi đọc UserLogin:", e)
    return CACHED_USER_LOGINS or []


@app.route('/api/login', methods=['POST'])
def api_login():
    """Hỗ trợ tương thích ngược đăng nhập nội bộ qua SharePoint UserLogin"""
    data = request.get_json(force=True, silent=True) or {}
    username = str(data.get('username', '')).strip()
    password = str(data.get('password', '')).strip()
    if not username or not password:
        return jsonify({'status': 'ERROR', 'message': 'Vui lòng nhập đầy đủ Tên đăng nhập và Mật khẩu!'}), 400

    users = fetch_user_logins(force=True)
    matched = None
    for u in users:
        if u.get('username', '').lower() == username.lower() and u.get('password') == password:
            matched = u
            break

    if not matched:
        return jsonify({'status': 'ERROR', 'message': 'Tên đăng nhập hoặc mật khẩu không chính xác!'}), 401

    roles = matched.get('adminRoles', [])
    session_token = secrets.token_urlsafe(32)
    session_data = {
        'username': matched.get('username'),
        'fullName': matched.get('fullName'),
        'department': matched.get('department'),
        'securityRole': matched.get('securityRole'),
        'adminRoles': roles,
        'loginAt': datetime.datetime.now(VN_TZ).strftime('%H:%M:%S %d/%m/%Y')
    }
    ACTIVE_SESSIONS[session_token] = {'user': session_data, 'createdAt': time.time()}
    return jsonify({
        'status': 'SUCCESS',
        'message': 'Đăng nhập thành công!',
        'token': session_token,
        'user': session_data
    })


@app.route('/api/auth/verify', methods=['GET', 'POST'])
def api_auth_verify():
    user = get_authenticated_user()
    if user:
        return jsonify({'status': 'SUCCESS', 'valid': True, 'user': user})
    # Nếu chưa có session UserLogin, kiểm tra tài khoản Microsoft
    token = get_access_token()
    acc = get_current_ms_account()
    if token and acc:
        return jsonify({
            'status': 'SUCCESS',
            'valid': True,
            'user': {
                'username': acc.get('username'),
                'fullName': acc.get('name'),
                'department': 'SeaZen Security Group',
                'securityRole': 'Microsoft 365 Account'
            }
        })
    return jsonify({'status': 'UNAUTHORIZED', 'valid': False, 'message': 'Phiên làm việc hết hạn hoặc chưa đăng nhập.'}), 401


@app.route('/api/auth/logout', methods=['POST'])
def api_auth_logout():
    token = request.headers.get('X-Auth-Token') or (request.get_json(force=True, silent=True) or {}).get('token', '')
    if token in ACTIVE_SESSIONS:
        del ACTIVE_SESSIONS[token]
    return jsonify({'status': 'SUCCESS', 'message': 'Đã đăng xuất'})


def open_browser():
    """Khởi động trình duyệt ở chế độ ứng dụng toàn màn hình"""
    time.sleep(1.2)
    edge_paths = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"
    ]
    opened = False
    for ep in edge_paths:
        if os.path.exists(ep):
            try:
                import subprocess
                subprocess.Popen([ep, "--app=http://localhost:5000", "--start-maximized"])
                opened = True
                break
            except Exception:
                pass
    if not opened:
        try:
            webbrowser.open("http://localhost:5000")
        except Exception:
            pass


if __name__ == '__main__':
    today = get_today_vn_str()
    lan_ip = get_lan_ip()
    hostname = socket.gethostname()
    current_acc = get_current_ms_account()

    print("===========================================================================")
    print("      SEAZEN CANTEEN CHECK-IN SYSTEM - STANDALONE KIOSK      ")
    print("===========================================================================")
    print(f"[*] Ngày làm việc (VN UTC+7) : {today}")
    print(f"[*] Thư mục ứng dụng         : {EXE_DIR}")
    print(f"[*] Thư mục AppData bảo mật  : {DATA_DIR}")
    print(f"[*] Lưu trữ Token DPAPI      : {TOKEN_CACHE_FILE}")
    if current_acc:
        print(f"[*] Tài khoản Microsoft 365  : {current_acc.get('username')} ({current_acc.get('name')})")
        print(f"[*] Trạng thái xác thực      : ĐÃ ĐĂNG NHẬP [HOẠT ĐỘNG]")
    else:
        print(f"[*] Tài khoản Microsoft 365  : CHƯA ĐĂNG NHẬP")
        print(f"[*] Trạng thái xác thực      : YÊU CẦU ĐĂNG NHẬP TRÊN GIAO DIỆN")
    print(f"[*] Tệp giao diện Kiosk      : {HTML_FILE}")
    print("---------------------------------------------------------------------------")
    print("   ĐỊA CHỈ TRUY CẬP HỆ THỐNG:")
    print("   👉 [1] Trên máy trạm Kiosk   : http://localhost:5000")
    print(f"   👉 [2] Từ mạng LAN (Theo IP) : http://{lan_ip}:5000")
    print(f"   👉 [3] Từ mạng LAN (Tên máy) : http://{hostname}:5000")
    print("---------------------------------------------------------------------------")
    print("[*] Sẵn sàng phục vụ điểm danh suất ăn Canteen SeaZen...")
    print("===========================================================================")

    threading.Thread(target=open_browser, daemon=True).start()
    app.run(host='0.0.0.0', port=5000, debug=False)
